"""Browser-independent adapters for the tested recruitment workflow."""
from __future__ import annotations
import copy
import hashlib
import importlib.util
import json
import re
import tempfile
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

STATUSES = ['demonstrated', 'partially_demonstrated', 'not_demonstrated', 'requires_human_confirmation']
BASE = Path(__file__).resolve().parent


def read_document(name, content):
    if len(content) > 10 * 1024 * 1024:
        raise ValueError('Upload a document smaller than 10 MB.')
    suffix = Path(name).suffix.lower()
    if suffix == '.txt':
        text = content.decode('utf-8-sig')
    elif suffix == '.pdf':
        from pypdf import PdfReader
        text = '\n'.join(page.extract_text() or '' for page in PdfReader(BytesIO(content)).pages)
    elif suffix == '.docx':
        from docx import Document
        document = Document(BytesIO(content))
        text = '\n'.join(p.text for p in document.paragraphs)
        for table in document.tables:
            text += '\n' + '\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)
    else:
        raise ValueError('Use a TXT, PDF or DOCX document.')
    if not text.strip():
        raise ValueError('No readable text found. Paste the text or use a text-based document.')
    return text.strip()


class RecruitmentEngine:
    def __init__(self, case_path, api_key=None):
        self.case_path = Path(case_path)
        spec = importlib.util.spec_from_file_location('recruitment_workflow_' + uuid.uuid4().hex, BASE / 'workflow.py')
        self.workflow = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(self.workflow)
        self.workflow.CASE_FILE = self.case_path
        self.workflow.save_case = self.save
        if api_key:
            from openai import OpenAI
            self.workflow.api_client = OpenAI(api_key=api_key, timeout=120.0, max_retries=0)

    def save(self, case, path=None):
        target = Path(path) if path else self.case_path
        target.parent.mkdir(parents=True, exist_ok=True)
        content = json.dumps(case, indent=2, ensure_ascii=False)
        # Keys and API clients are never members of the case dictionary.
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=target.parent, suffix='.tmp', delete=False) as f:
            f.write(content)
            temporary = Path(f.name)
        try:
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)

    def load(self):
        return json.loads(self.case_path.read_text(encoding='utf-8'))

    def require_client(self):
        if not hasattr(self.workflow, 'api_client'):
            raise ValueError('Enter an API key in Settings before generating a draft.')

    @staticmethod
    def confirm(reviewer, confirmed):
        if not reviewer.strip() or not confirmed:
            raise ValueError('Enter your reviewer name and tick the confirmation.')

    def generate_role(self, brief):
        self.require_client()
        if not brief.strip() or len(brief) > 16000:
            raise ValueError('Provide a hiring brief of 1–16,000 characters.')
        w = self.workflow
        record, extraction = w.call_model(w.extraction_prompt, 'HIRING BRIEF — SOURCE DATA:\n' + brief, w.brief_schema, lambda r: w.check_source_quotes(r, brief))
        materials, drafting = w.call_model(w.drafting_prompt, json.dumps({'original_client_brief': brief, 'extracted_hiring_record': record}, ensure_ascii=False), w.materials_schema, lambda r: w.check_scorecard(r, record))
        case = {'source_brief': brief, 'hiring_record': record, 'recruitment_materials': materials,
                'traces': {'extraction': extraction, 'drafting': drafting}, 'criteria_approved': False,
                'status': 'draft_requires_human_review'}
        self.save(case)
        return case

    def approve_role(self, case, materials, reviewer, confirmed):
        self.confirm(reviewer, confirmed)
        w = self.workflow
        review = case.get('review', {})
        if review.get('approval') and materials == review.get('materials'):
            # Verify the approval before preserving it. Reconfirming unchanged
            # materials must not create a new event or invalidate screening.
            w.get_approved_case(case)
            return
        w.begin_review(case)
        w.save_review(case, materials, w.review_source_version(case))
        w.approve_review(case, reviewer, confirmed)
        self.save(case)

    def screen(self, case, candidate_id, text, confirmed):
        self.require_client()
        if not confirmed:
            raise ValueError('Check the CV text and confirm it before screening.')
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,64}', candidate_id):
            raise ValueError('Use letters, numbers, underscores or hyphens for the candidate ID.')
        w = self.workflow
        approved = w.get_approved_case(case)
        old = case.get('screening_results', {}).get(candidate_id)
        if old and old.get('candidate_source') == text and old.get('approval') == approved['approval']:
            return old
        if case.get('screening_reviews', {}).get(candidate_id, {}).get('approval'):
            raise PermissionError('This candidate already has approved screening. Use a new candidate ID for a new CV version.')
        result = w.screen_candidate(case, candidate_id, text)
        self.save(case)
        return result

    def approve_screening(self, case, candidate_id, edited_rows, reviewer, confirmed):
        self.confirm(reviewer, confirmed)
        w = self.workflow
        version = w.screening_review_version(case, candidate_id)
        draft = case['screening_results'][candidate_id]
        expected = {r['criterion_id']: r for r in draft['assessments']}
        evidence = w.candidate_evidence(candidate_id, draft['candidate_source'])
        self.check_rows(edited_rows, expected)
        result = {'assessments': {r['criterion_id']: {'status': r['status'], 'evidence_ids': r['evidence_ids'], 'explanation': r['explanation']} for r in edited_rows}}
        w.validate_screening(result, w.screening_schema(expected), evidence)
        source = {r['id']: r['text'] for r in evidence['evidence']}
        rows = [{**copy.deepcopy(expected[r['criterion_id']]), 'status': r['status'], 'explanation': r['explanation'],
                 'evidence': [{'id': ref, 'quote': source[ref]} for ref in r['evidence_ids']]} for r in edited_rows]
        case.setdefault('screening_reviews', {})[candidate_id] = {
            'draft_version': version, 'assessments': rows,
            'approval': {'reviewer': reviewer.strip(), 'approved_at_utc': datetime.now(timezone.utc).isoformat(),
                         'draft_version': version, 'rows_version': w.review_fingerprint(rows)}}
        self.save(case)

    @staticmethod
    def check_rows(rows, expected):
        ids = [r['criterion_id'] for r in rows]
        if len(ids) != len(expected) or set(ids) != set(expected):
            raise ValueError('Keep every criterion exactly once.')
        for row in rows:
            if row['criterion'] != expected[row['criterion_id']]['criterion']:
                raise ValueError('Criterion wording must match the approved scorecard.')

    def plan(self, case, candidate_id):
        self.require_client()
        w = self.workflow
        context = w.read_interview_context(case, candidate_id)
        existing = case.get('interview_plans', {}).get(candidate_id)
        if existing and existing.get('context_version') == w.review_fingerprint(context):
            if existing.get('approval'):
                w.get_approved_interview_plan(case, candidate_id)
            return existing
        if existing and existing.get('approval'):
            raise PermissionError('An approved plan has outdated sources. Review the changed source records first.')
        return w.plan_interview(case, candidate_id)

    def approve_plan(self, case, candidate_id, rows, notes, reviewer, confirmed):
        self.confirm(reviewer, confirmed)
        w = self.workflow
        context = w.read_interview_context(case, candidate_id)
        plan = case['interview_plans'][candidate_id]
        if plan['context_version'] != w.review_fingerprint(context):
            raise PermissionError('Source records changed. Generate a new plan.')
        self.check_rows(rows, context['approved_criteria'])
        result = {'plan': {r['criterion_id']: {'reason': r['reason'], 'questions': r['questions']} for r in rows}, 'interviewer_notes': notes}
        w.validate_interview_plan(result, w.interview_plan_schema(context['approved_criteria']))
        plan['plan'] = copy.deepcopy(rows)
        plan['interviewer_notes'] = notes
        plan['approval'] = {'reviewer': reviewer.strip(), 'approved_at': datetime.now(timezone.utc).isoformat(),
                            'context_version': plan['context_version'], 'plan_version': w.review_fingerprint({'plan': rows, 'interviewer_notes': notes})}
        plan['status'] = 'interview_plan_approved'
        self.save(case)

    def save_notes(self, case, candidate_id, text, source_type, reviewer, confirmed):
        self.confirm(reviewer, confirmed)
        if not text.strip() or len(text.strip()) > 16000:
            raise ValueError('Provide interview notes of 1–16,000 characters.')
        if source_type not in {'actual', 'simulated'}:
            raise ValueError('Choose actual or simulated interview evidence.')
        w = self.workflow
        plan = w.get_approved_interview_plan(case, candidate_id)
        text = text.strip()
        case.setdefault('interview_records', {})[candidate_id] = {
            'candidate_id': candidate_id, 'source_type': source_type, 'source_text': text,
            'evidence': [{'evidence_id': f'I{i:03d}', 'source_quote': line.strip()} for i, line in enumerate([l for l in text.splitlines() if l.strip()], 1)],
            'interview_plan_version': plan['approval']['plan_version'],
            'review': {'reviewer': reviewer.strip(), 'reviewed_at': datetime.now(timezone.utc).isoformat(),
                       'notes_version': w.review_fingerprint({'source_text': text, 'source_type': source_type})}}
        self.save(case)

    def evaluate(self, case, candidate_id):
        self.require_client()
        return self.workflow.evaluate_interview(case, candidate_id)

    def approve_evaluation(self, case, candidate_id, edited_rows, summary, next_step, reviewer, confirmed):
        self.confirm(reviewer, confirmed)
        w = self.workflow
        context = w.read_evaluation_context(case, candidate_id)
        draft = case['interview_evaluations'][candidate_id]
        if draft['context_version'] != w.review_fingerprint(context):
            raise PermissionError('Source records changed. This evaluation is outdated.')
        expected = {r['criterion_id']: r for r in draft['assessments']}
        self.check_rows(edited_rows, expected)
        result = {'assessments': {r['criterion_id']: {'status': r['status'], 'basis': 'insufficient_information',
                  'evidence_ids': r['evidence_ids'], 'explanation': r['explanation'], 'open_questions': r['open_questions']} for r in edited_rows},
                  'client_summary': summary.strip(), 'proposed_next_step': next_step.strip()}
        w.validate_evaluation(result, w.evaluation_schema(context['approved_criteria']), context)
        source = {r['id']: r for r in context['evidence']}
        for key, row in result['assessments'].items():
            if not set(re.findall(r'\b[EI]\d{3,}\b', row['explanation'])).issubset(set(row['evidence_ids'])):
                raise ValueError(key + ': select all evidence IDs mentioned in the explanation.')
        review = {'candidate_id': candidate_id, 'draft_version': w.review_fingerprint(draft), 'context_version': draft['context_version'],
                  'assessments': [{'criterion_id': key, 'criterion': expected[key]['criterion'], 'status': row['status'], 'basis': row['basis'],
                                   'evidence': [copy.deepcopy(source[ref]) for ref in row['evidence_ids']], 'explanation': row['explanation'],
                                   'open_questions': row['open_questions']} for key, row in result['assessments'].items()],
                  'client_summary': result['client_summary'], 'proposed_next_step': result['proposed_next_step'], 'approval': None}
        review['approval'] = {'reviewer': reviewer.strip(), 'approved_at': datetime.now(timezone.utc).isoformat(),
                              'report_version': w.review_fingerprint(w.evaluation_review_payload(review))}
        case.setdefault('evaluation_reviews', {})[candidate_id] = review
        self.save(case)

    def tracker(self, case, candidate_id, stage, action, reviewer, confirmed):
        return self.workflow.update_recruitment_tracker(case, candidate_id, stage, action, reviewer, confirmed)

    def quality(self, case, candidate_id):
        w = self.workflow
        context = w.read_evaluation_context(case, candidate_id)
        review = w.get_approved_evaluation(case, candidate_id)
        expected = set(context['approved_criteria'])
        source = {r['id']: r for r in context['evidence']}
        draft = {r['criterion_id']: r for r in case['interview_evaluations'][candidate_id]['assessments']}
        quotes = [q for r in review['assessments'] for q in r['evidence']]
        changes = [{'criterion_id': r['criterion_id'], 'model_status': draft[r['criterion_id']]['status'], 'reviewed_status': r['status']}
                   for r in review['assessments'] if r['status'] != draft[r['criterion_id']]['status']]
        return {'criterion_coverage_passed': len(review['assessments']) == len(expected) and {r['criterion_id'] for r in review['assessments']} == expected,
                'criteria_covered': len(review['assessments']), 'criteria_expected': len(expected),
                'exact_source_quotations': sum(q == source.get(q['id']) for q in quotes), 'quotations_checked': len(quotes), 'status_changes': changes}

    def report_markdown(self, case, candidate_id):
        review = self.workflow.get_approved_evaluation(case, candidate_id)
        lines = [f'# Recruitment report — {candidate_id}', '', review['client_summary'], '']
        for row in review['assessments']:
            lines += [f"## {row['criterion']}", '', f"Status: {row['status']}", '', row['explanation'], '']
            for e in row['evidence']:
                lines += [f"- {e['id']} [{e['source']}]: {e['text']}"]
            for q in row['open_questions']:
                lines += [f'- Follow-up: {q}']
            lines += ['']
        lines += ['## Proposed next step', '', review['proposed_next_step'], '', f"Reviewed by: {review['approval']['reviewer']}"]
        return '\n'.join(lines)

    def safeguards(self, case, candidate_id):
        """Exercise gates on copies; never save or alter the user's case."""
        w = self.workflow
        before = w.review_fingerprint(case)
        results = []
        try:
            w.get_approved_evaluation(copy.deepcopy(case), candidate_id)
            results.append({'check': 'Unchanged approved report accepted', 'passed': True})
        except (ValueError, PermissionError):
            results.append({'check': 'Unchanged approved report accepted', 'passed': False})
        tests = [
            ('Missing role approval', lambda c: c['review'].update(approval=None)),
            ('Criteria changed after approval', lambda c: c['review']['materials']['draft_scorecard'][0].update(criterion='Changed test criterion')),
            ('CV changed after screening review', lambda c: c['screening_results'][candidate_id].update(candidate_source='Changed test CV')),
            ('Interview plan changed after approval', lambda c: c['interview_plans'][candidate_id]['plan'][0].update(reason='Changed test reason')),
            ('Interview notes changed after review', lambda c: c['interview_records'][candidate_id].update(source_text='Changed test notes')),
            ('Final report changed after approval', lambda c: c['evaluation_reviews'][candidate_id]['assessments'][0].update(explanation='Changed test explanation'))]
        for label, modify in tests:
            copied = copy.deepcopy(case)
            modify(copied)
            try:
                w.get_approved_evaluation(copied, candidate_id)
                passed = False
            except (ValueError, PermissionError):
                passed = True
            results.append({'check': label, 'passed': passed})
        # Confirmation is checked before any tracker write in the tested function.
        try:
            w.update_recruitment_tracker(copy.deepcopy(case), candidate_id, 'follow_up_required', 'Test only', 'Test reviewer', False)
            passed = False
        except PermissionError:
            passed = True
        results.append({'check': 'Tracker update without human confirmation', 'passed': passed})
        if w.review_fingerprint(case) != before:
            raise RuntimeError('Safeguard test altered the original case.')
        return {'passed': sum(r['passed'] for r in results), 'expected': len(results), 'original_case_unchanged': True, 'checks': results}
